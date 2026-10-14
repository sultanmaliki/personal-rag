from ingest import document_ingest as di


def test_walk_documents_reads_supported_text_formats(tmp_path):
    (tmp_path / "notes.txt").write_text("plain text notes", encoding="utf-8")
    (tmp_path / "resume.md").write_text("# Resume\nsome markdown", encoding="utf-8")

    found = dict(di.walk_documents(tmp_path))
    assert found["notes.txt"] == "plain text notes"
    assert found["resume.md"] == "# Resume\nsome markdown"


def test_walk_documents_skips_unsupported_extensions(tmp_path):
    (tmp_path / "notes.txt").write_text("keep me", encoding="utf-8")
    (tmp_path / "archive.zip").write_bytes(b"PK\x03\x04junk")
    (tmp_path / "image.png").write_bytes(b"\x89PNG\r\n\x1a\n")

    found = dict(di.walk_documents(tmp_path))
    assert set(found) == {"notes.txt"}


def test_walk_documents_skips_empty_and_oversized_files(tmp_path, monkeypatch):
    (tmp_path / "empty.txt").write_text("", encoding="utf-8")
    (tmp_path / "normal.txt").write_text("real content", encoding="utf-8")

    monkeypatch.setattr(di, "MAX_FILE_SIZE", 20)
    (tmp_path / "toobig.txt").write_text("way more than twenty bytes of content", encoding="utf-8")

    found = dict(di.walk_documents(tmp_path))
    assert "empty.txt" not in found
    assert "toobig.txt" not in found
    assert found["normal.txt"] == "real content"


def test_walk_documents_recurses_into_subfolders(tmp_path):
    sub = tmp_path / "subfolder"
    sub.mkdir()
    (sub / "nested.txt").write_text("nested content", encoding="utf-8")

    found = dict(di.walk_documents(tmp_path))
    assert found["subfolder/nested.txt"] == "nested content"


def test_walk_documents_skips_file_that_fails_extraction(tmp_path, monkeypatch):
    """A single corrupt/malformed document shouldn't abort the whole run."""
    (tmp_path / "good.txt").write_text("fine", encoding="utf-8")
    (tmp_path / "bad.pdf").write_bytes(b"not actually a pdf")

    found = dict(di.walk_documents(tmp_path))
    assert "good.txt" in found
    assert "bad.pdf" not in found


def test_extract_text_dispatches_by_suffix(tmp_path, monkeypatch):
    calls = []
    monkeypatch.setattr(di, "_extract_pdf", lambda p: calls.append(("pdf", p)) or "pdf text")
    monkeypatch.setattr(di, "_extract_docx", lambda p: calls.append(("docx", p)) or "docx text")

    pdf_path = tmp_path / "a.pdf"
    docx_path = tmp_path / "b.docx"
    txt_path = tmp_path / "c.txt"
    txt_path.write_text("plain", encoding="utf-8")

    assert di._extract_text(pdf_path) == "pdf text"
    assert di._extract_text(docx_path) == "docx text"
    assert di._extract_text(txt_path) == "plain"
    assert calls == [("pdf", pdf_path), ("docx", docx_path)]


def test_extract_docx_reads_paragraphs_and_tables(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("Hello from a real docx fixture.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "left"
    table.rows[0].cells[1].text = "right"
    path = tmp_path / "fixture.docx"
    doc.save(str(path))

    text = di._extract_docx(path)
    assert "Hello from a real docx fixture." in text
    assert "left | right" in text


def test_purge_stale_documents_removes_file_no_longer_present(monkeypatch):
    monkeypatch.setattr(
        di.vectorstore,
        "list_metadatas",
        lambda where: [{"path": "gone.txt"}, {"path": "kept.txt"}],
    )
    deleted = []
    monkeypatch.setattr(di.vectorstore, "delete_where", lambda where: deleted.append(where))

    di._purge_stale_documents({"kept.txt"})

    assert deleted == [{"source": "document", "path": "gone.txt"}]


def test_ingest_end_to_end_with_isolated_config(monkeypatch, tmp_path):
    (tmp_path / "a.txt").write_text("content a", encoding="utf-8")
    (tmp_path / "b.md").write_text("content b", encoding="utf-8")

    import types

    fake_config = types.SimpleNamespace(info_dir=tmp_path)
    monkeypatch.setattr(di, "config", fake_config)
    monkeypatch.setattr(di.vectorstore, "list_metadatas", lambda where: [])
    monkeypatch.setattr(di.vectorstore, "delete_where", lambda where: None)

    stored_chunks = []
    monkeypatch.setattr(di, "embed_and_store", lambda chunks, label="chunks": stored_chunks.extend(chunks))

    di.ingest()

    stored_paths = {c.metadata["path"] for c in stored_chunks}
    assert stored_paths == {"a.txt", "b.md"}
    assert all(c.metadata["source"] == "document" for c in stored_chunks)
