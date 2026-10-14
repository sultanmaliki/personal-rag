"""Ingest local documents dropped into info/ (PDF, Word, plain text,
Markdown) so personal notes, resumes, contracts, etc. become part of the
RAG knowledge base alongside GitHub repos and the crawled website.
"""
from __future__ import annotations

from pathlib import Path

from app.config import config
from app.rag import vectorstore
from ingest.common import chunks_from_document, embed_and_store

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}
MAX_FILE_SIZE = 20_000_000  # 20MB -- generous for documents, still bounded


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n\n".join(pages)


def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    # .txt, .md: read directly
    return path.read_text(encoding="utf-8", errors="ignore")


def walk_documents(info_dir: Path):
    """Yields (relative_path, text) for every supported file under info_dir."""
    for path in sorted(info_dir.rglob("*")):
        if path.is_dir():
            continue
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        try:
            size = path.stat().st_size
        except OSError:
            continue
        if size == 0 or size > MAX_FILE_SIZE:
            continue
        try:
            text = _extract_text(path)
        except Exception as exc:  # a malformed/corrupt file shouldn't abort the whole run
            print(f"  skip {path.name} (failed to extract text: {exc})")
            continue
        text = text.strip()
        if text:
            yield path.relative_to(info_dir).as_posix(), text


def _purge_stale_documents(kept_paths: set[str]) -> None:
    stored = {m["path"] for m in vectorstore.list_metadatas({"source": "document"}) if m.get("path")}
    for stale in stored - kept_paths:
        print(f"purging stale document from index: {stale}")
        vectorstore.delete_where({"source": "document", "path": stale})


def ingest() -> None:
    info_dir = config.info_dir
    info_dir.mkdir(parents=True, exist_ok=True)

    print(f"Scanning {info_dir} for documents ({', '.join(sorted(SUPPORTED_SUFFIXES))})...")
    kept_paths: set[str] = set()
    for rel_path, text in walk_documents(info_dir):
        kept_paths.add(rel_path)
        metadata = {"source": "document", "path": rel_path, "title": Path(rel_path).stem}
        doc_chunks = chunks_from_document(f"document:{rel_path}", text, metadata)
        # Clear this document's previous chunks first so a shrunk/edited
        # file doesn't leave stale chunks behind alongside the fresh ones.
        vectorstore.delete_where({"source": "document", "path": rel_path})
        print(f"  {rel_path} -> {len(doc_chunks)} chunks")
        embed_and_store(doc_chunks, label=f"chunks ({rel_path})")

    _purge_stale_documents(kept_paths)

    if not kept_paths:
        print(f"  no supported documents found in {info_dir} -- drop a .pdf/.docx/.txt/.md file there and re-run")


if __name__ == "__main__":
    ingest()
